import os
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional
import fitz  # PyMuPDF
import docx
import openpyxl

from app.core.config import settings

@dataclass
class FileRecord:
    name: str
    path: str
    size_bytes: int
    extension: str
    content: Optional[str]   # None if binary/too large/unreadable

@dataclass
class FolderIndex:
    base_folder: str
    files: list[FileRecord] = field(default_factory=list)
    total_files: int = 0
    indexed_at: str = ""


def extract_text(file_path: Path, ext: str, max_mb: int) -> Optional[str]:
    size_mb = file_path.stat().st_size / (1024 * 1024)
    if size_mb > max_mb:
        return None  # too large to index content

    try:
        if ext in [".txt", ".md", ".csv", ".log", ".json", ".xml", ".yaml", ".yml", ".py", ".js", ".ts", ".html", ".css"]:
            return file_path.read_text(encoding="utf-8", errors="ignore")


        elif ext == ".pdf":
            doc = fitz.open(str(file_path))
            text = []
            for page in doc:
                t = page.get_text("text")
                if t.strip():
                    text.append(t)
            return "\n".join(text) if text else None

        elif ext == ".docx":
            doc = docx.Document(str(file_path))
            return "\n".join(p.text for p in doc.paragraphs)

        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    lines.append(" | ".join(str(c) for c in row if c is not None))
            return "\n".join(lines)


    except Exception as e:
        print(f"[ERROR] Failed to read {file_path}: {e}")

    return None  # unreadable binary or unsupported format


def build_index(folder_path: str) -> FolderIndex:
    from datetime import datetime

    base = Path(folder_path).resolve()
    index = FolderIndex(base_folder=str(base), indexed_at=datetime.now().isoformat())
    max_mb = settings.max_file_size_mb

    for root, _, files in os.walk(base):
        for filename in files:
            fp = Path(root) / filename
            ext = fp.suffix.lower()
            try:
                size = fp.stat().st_size
                content = extract_text(fp, ext, max_mb)
                index.files.append(FileRecord(
                    name=filename,
                    path=str(fp),
                    size_bytes=size,
                    extension=ext,
                    content=content
                ))
            except Exception as e:
                print(f"[INDEX ERROR] {fp}: {e}")

    index.total_files = len(index.files)
    return index